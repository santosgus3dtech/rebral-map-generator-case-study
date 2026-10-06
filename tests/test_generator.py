import json
from pathlib import Path

from openpyxl import load_workbook

from map_demo.generator import generate_workbook


ROOT = Path(__file__).resolve().parents[1]


def test_generator_preserves_structure_and_dynamic_total(tmp_path):
    payload = json.loads((ROOT / "examples" / "synthetic-map.json").read_text(encoding="utf-8"))
    output = generate_workbook(payload, tmp_path / "map.xlsx")
    workbook = load_workbook(output, data_only=False)
    assert workbook.sheetnames == ["Expense Map", "Validation"]
    sheet = workbook["Expense Map"]
    assert sheet["A1"].value == "SYNTHETIC EXPENSE MAP"
    assert sheet["F14"].value == "=SUM(F9:F13)"
    assert sheet.auto_filter.ref == "A8:F13"
    assert "#REF!" not in " ".join(str(cell.value) for row in sheet.iter_rows() for cell in row)


def test_generator_rejects_negative_amount(tmp_path):
    payload = json.loads((ROOT / "examples" / "synthetic-map.json").read_text(encoding="utf-8"))
    payload["documents"][0]["amount"] = -1
    try:
        generate_workbook(payload, tmp_path / "map.xlsx")
    except ValueError as error:
        assert "negative" in str(error)
    else:
        raise AssertionError("negative amount should be rejected")
